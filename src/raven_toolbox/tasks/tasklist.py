"""Parse a metabolic task list.

A task list defines, per task, allowed inputs and outputs, optional extra reactions
(equations), reaction-bound changes, and whether the task should fail. Tasks
are checked with :func:`raven_toolbox.tasks.check_tasks`.

The file is tab-delimited (for example ``.txt`` or ``.tsv``) or Excel (``.xlsx`` or
``.xlsm``, sheet ``TASKS``; needs the ``[excel]`` extra). Recognised columns (the only
required one is ``ID``, which must not be the first column; the first column is reserved
for comment markers, and a row with text in the first column is skipped, as in RAVEN's
``parseTaskList``):

    ID · DESCRIPTION · IN · IN LB · IN UB · OUT · OUT LB · OUT UB ·
    EQU · EQU LB · EQU UB · CHANGED RXN · CHANGED LB · CHANGED UB ·
    SHOULD FAIL · PRINT FLUX · COMMENTS

A task spans consecutive rows; only its first row carries an ID. Metabolites are
written ``name[compartment]``; several in one cell are separated by ``;`` (sharing
that row's bounds). ``IN`` and ``OUT`` default to LB 0 and UB 1000; ``EQU`` defaults to
LB -1000 if reversible (``<=>``) else 0, and UB 1000. An empty ``CHANGED LB`` or
``CHANGED UB`` leaves that bound of the reaction unchanged. ``SHOULD FAIL`` and
``PRINT FLUX`` are false when empty or ``0``, ``false`` or ``no`` (any case), true
otherwise. The special tokens ``ALLMETS`` and ``ALLMETSIN[comp]`` allow free uptake or
excretion of all metabolites (only the upper bound is used). A row whose ID starts with
``#`` is skipped.
"""
from __future__ import annotations

import csv
import warnings
from dataclasses import dataclass, field
from pathlib import Path

_COLUMNS = (
    "ID", "DESCRIPTION", "IN", "IN LB", "IN UB", "OUT", "OUT LB", "OUT UB",
    "EQU", "EQU LB", "EQU UB", "CHANGED RXN", "CHANGED LB", "CHANGED UB",
    "SHOULD FAIL", "PRINT FLUX", "COMMENTS",
)


@dataclass
class Task:
    """One metabolic task. Bounds are ``(metabolite_or_reaction, lb, ub)`` triples.

    In ``changed``, a bound is ``None`` when the cell was empty, meaning "leave that
    bound of the reaction as it is" (RAVEN reads it as NaN).
    """

    id: str
    description: str = ""
    should_fail: bool = False
    print_fluxes: bool = False
    comments: str = ""
    inputs: list[tuple[str, float, float]] = field(default_factory=list)
    outputs: list[tuple[str, float, float]] = field(default_factory=list)
    equations: list[tuple[str, float, float]] = field(default_factory=list)
    changed: list[tuple[str, float | None, float | None]] = field(default_factory=list)


def _truthy(value: str) -> bool:
    return value.strip().lower() not in ("", "0", "false", "no")


def _num(value: str, default: float) -> float:
    value = value.strip()
    return float(value) if value else default


def _opt_num(value: str) -> float | None:
    value = value.strip()
    return float(value) if value else None


def _read_rows(path: str | Path) -> list[list[str]]:
    path = Path(path)
    if path.suffix.lower() == ".xls":
        raise ValueError(
            f"{path}: the old binary .xls format is not supported. Save the task list as "
            ".xlsx (sheet 'TASKS') or as a tab-delimited .txt file."
        )
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - optional dep
            raise ImportError("Reading .xlsx task lists needs the '[excel]' extra (openpyxl).") from exc
        wb = load_workbook(path, data_only=True)
        if "TASKS" not in wb.sheetnames:
            raise ValueError(
                f"{path}: workbook has no sheet named 'TASKS' "
                f"(found: {wb.sheetnames}). Rename the sheet or pick that file."
            )
        ws = wb["TASKS"]
        return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
    with open(path, encoding="utf-8", newline="") as handle:
        return [row for row in csv.reader(handle, delimiter="\t")]


def parse_task_list(path: str | Path) -> list[Task]:
    """Parse a task-list file into :class:`Task` objects."""
    rows = _read_rows(path)
    header_idx = next(
        (i for i, r in enumerate(rows) if any(c.strip().upper() == "ID" for c in r)), None
    )
    if header_idx is None:
        raise ValueError(f"{path}: no header row with an 'ID' column found.")
    header = [c.strip().upper() for c in rows[header_idx]]
    col = {name: header.index(name) for name in _COLUMNS if name in header}
    if col["ID"] == 0:
        raise ValueError(
            f"{path}: the TASKS sheet must have a column named ID that is not the first "
            "column; the first column is reserved for comment markers."
        )

    def cell(row: list[str], name: str) -> str:
        i = col.get(name)
        return row[i].strip() if i is not None and i < len(row) else ""

    # Columns whose presence on an ID-less row signals real continuation data
    # (vs. pure whitespace/comment), used by the orphan-row warning below.
    _DATA_COLS = ("IN", "OUT", "EQU", "CHANGED RXN")

    tasks: list[Task] = []
    current: Task | None = None
    for row_no, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(c.strip() for c in row):
            continue
        # RAVEN parseTaskList: "all rows starting with a non-empty [first] cell are
        # removed". Comment / section-header rows carry a marker (e.g. "#") in the first
        # column, with the ID column holding descriptive text; without this they become
        # spurious empty tasks (69 vs RAVEN's 57 on metabolicTasks_Essential).
        if row and row[0].strip():
            continue
        rid = cell(row, "ID")
        if rid.startswith("#"):
            continue
        if rid:
            current = Task(
                id=rid,
                description=cell(row, "DESCRIPTION"),
                should_fail=_truthy(cell(row, "SHOULD FAIL")),
                print_fluxes=_truthy(cell(row, "PRINT FLUX")),
                comments=cell(row, "COMMENTS"),
            )
            tasks.append(current)
        if current is None:
            # Continuation row appearing before any task ID: warn and skip rather than
            # silently dropping it, which would mask malformed task files.
            if any(cell(row, c) for c in _DATA_COLS):
                warnings.warn(
                    f"{path}: row {row_no} carries task data but no task ID has "
                    "been seen yet; the row is being skipped.",
                    stacklevel=2,
                )
            continue
        _add_row(current, row, cell)
    return tasks


def _add_row(task: Task, row: list[str], cell) -> None:
    if inp := cell(row, "IN"):
        lb, ub = _num(cell(row, "IN LB"), 0.0), _num(cell(row, "IN UB"), 1000.0)
        task.inputs += [(m.strip(), lb, ub) for m in inp.split(";") if m.strip()]
    if out := cell(row, "OUT"):
        lb, ub = _num(cell(row, "OUT LB"), 0.0), _num(cell(row, "OUT UB"), 1000.0)
        task.outputs += [(m.strip(), lb, ub) for m in out.split(";") if m.strip()]
    if equ := cell(row, "EQU"):
        lb = _num(cell(row, "EQU LB"), -1000.0 if "<=>" in equ else 0.0)
        ub = _num(cell(row, "EQU UB"), 1000.0)
        task.equations.append((equ.strip(), lb, ub))
    if chg := cell(row, "CHANGED RXN"):
        chg_lb, chg_ub = _opt_num(cell(row, "CHANGED LB")), _opt_num(cell(row, "CHANGED UB"))
        task.changed += [(r.strip(), chg_lb, chg_ub) for r in chg.split(";") if r.strip()]
